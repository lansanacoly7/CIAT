"""Tests unitaires et d'intégration pour le module d'analyse (apps.analytics).

Vérifie :
1. Contrôle d'accès strict (RBAC) :
   - Coordonnateur, Administrateur, Membre du comité -> 200 OK
   - Porteur, Visiteur, Mentor, Anonyme -> 403 Forbidden / 401 Unauthorized
2. Conformité stricte aux contrats d'API (CONTRATS_API.md, section P3)
3. Précision des agrégats (idées reçues, dossiers déposés, taux de sélection, etc.)
4. Combinaison dynamique des filtres (?debut=&fin=&cohorte=&domaine=&statut=)
5. Robustesse aux cas limites (base vide, 0 division)
"""
from datetime import date, timedelta
from decimal import Decimal
import io
from unittest.mock import patch
import openpyxl

from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APIClient

from apps.core.models import Cohorte, Domaine, Projet, Role, User
from apps.analytics.models import (
    Financement,
    Partenariat,
    StatutFinancement,
    StatutPartenariat,
    TypePartenariat,
)


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class AnalyticsSecurityAndRBACPermissionsTests(TestCase):
    """Vérification stricte des autorisations RBAC sur tous les endpoints d'analyse."""

    def setUp(self):
        self.client = APIClient()

        # Utilisateurs avec différents rôles
        self.admin = User.objects.create_user(
            email="admin.test@ciat.sn",
            password="pwd",
            role=Role.ADMIN,
        )
        self.coordo = User.objects.create_user(
            email="coordo.test@ciat.sn",
            password="pwd",
            role=Role.COORDONNATEUR,
        )
        self.comite = User.objects.create_user(
            email="comite.test@ciat.sn",
            password="pwd",
            role=Role.COMITE,
        )
        self.porteur = User.objects.create_user(
            email="porteur.test@ciat.sn",
            password="pwd",
            role=Role.PORTEUR,
        )
        self.visiteur = User.objects.create_user(
            email="visiteur.test@ciat.sn",
            password="pwd",
            role=Role.VISITEUR,
        )
        self.mentor = User.objects.create_user(
            email="mentor.test@ciat.sn",
            password="pwd",
            role=Role.MENTOR,
        )

        self.endpoints = [
            reverse("analytics-indicateurs"),
            reverse("analytics-par-domaine"),
            reverse("analytics-par-cohorte"),
            reverse("analytics-par-genre"),
            reverse("analytics-evolution"),
        ]


    def test_anonymous_access_is_denied(self):
        """Un utilisateur non authentifié ne peut accéder à aucun endpoint d'analyse."""
        for url in self.endpoints:
            res = self.client.get(url)
            self.assertIn(
                res.status_code,
                [status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN],
                f"L'URL {url} doit être protégée contre les requêtes anonymes.",
            )

    def test_porteur_is_forbidden_403(self):
        """Un porteur de projet reçoit une erreur 403 Forbidden sur tous les endpoints."""
        self.client.force_authenticate(user=self.porteur)
        for url in self.endpoints:
            res = self.client.get(url)
            self.assertEqual(
                res.status_code,
                status.HTTP_403_FORBIDDEN,
                f"Le porteur ne doit pas accéder à {url}",
            )

    def test_visiteur_and_mentor_are_forbidden_403(self):
        """Visiteurs et mentors ne peuvent pas accéder au tableau de bord décisionnel."""
        for user in [self.visiteur, self.mentor]:
            self.client.force_authenticate(user=user)
            for url in self.endpoints:
                res = self.client.get(url)
                self.assertEqual(
                    res.status_code,
                    status.HTTP_403_FORBIDDEN,
                    f"L'utilisateur {user.role} ne doit pas accéder à {url}",
                )

    def test_authorized_roles_can_access_200(self):
        """Coordonnateur, Administrateur et Membre du comité accèdent avec succès (200 OK)."""
        for user in [self.coordo, self.admin, self.comite]:
            self.client.force_authenticate(user=user)
            for url in self.endpoints:
                res = self.client.get(url)
                self.assertEqual(
                    res.status_code,
                    status.HTTP_200_OK,
                    f"L'utilisateur {user.role} doit pouvoir accéder à {url}",
                )


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class AnalyticsEndpointsAndDataTests(TestCase):
    """Vérification des calculs, formats de réponse et filtres."""

    def setUp(self):
        self.client = APIClient()
        self.coordo = User.objects.create_user(
            email="coordo.data@ciat.sn",
            password="pwd",
            role=Role.COORDONNATEUR,
        )
        self.client.force_authenticate(user=self.coordo)

        # Création des porteurs (genres variés)
        self.porteur_femme = User.objects.create_user(
            email="femme@ciat.sn",
            password="pwd",
            role=Role.PORTEUR,
            genre=User.Genre.FEMME,
        )
        self.porteur_homme = User.objects.create_user(
            email="homme@ciat.sn",
            password="pwd",
            role=Role.PORTEUR,
            genre=User.Genre.HOMME,
        )
        self.porteur_autre = User.objects.create_user(
            email="autre@ciat.sn",
            password="pwd",
            role=Role.PORTEUR,
            genre=User.Genre.AUTRE,
        )

        # Domaines
        self.domaine_tic = Domaine.objects.create(nom="TIC et Numérique", slug="tic-numerique")
        self.domaine_agri = Domaine.objects.create(nom="Agriculture", slug="agriculture")

        # Cohortes
        self.cohorte_2024 = Cohorte.objects.create(nom="Cohorte 2024", annee=2024)
        self.cohorte_2025 = Cohorte.objects.create(nom="Cohorte 2025", annee=2025)

        # Création de projets aux statuts variés
        # 1. Idée
        Projet.objects.create(
            titre="Projet Idée TIC",
            domaine=self.domaine_tic,
            cohorte=self.cohorte_2024,
            porteur=self.porteur_femme,
            statut=Projet.Statut.IDEE,
        )
        # 2. Déposé
        Projet.objects.create(
            titre="Projet Déposé Agri",
            domaine=self.domaine_agri,
            cohorte=self.cohorte_2024,
            porteur=self.porteur_homme,
            statut=Projet.Statut.DEPOSE,
        )
        # 3. Sélectionné
        Projet.objects.create(
            titre="Projet Sélectionné TIC",
            domaine=self.domaine_tic,
            cohorte=self.cohorte_2025,
            porteur=self.porteur_femme,
            statut=Projet.Statut.SELECTIONNE,
        )
        # 4. Incubé
        Projet.objects.create(
            titre="Projet Incubé TIC",
            domaine=self.domaine_tic,
            cohorte=self.cohorte_2025,
            porteur=self.porteur_autre,
            statut=Projet.Statut.INCUBE,
        )
        # 5. Prototype
        Projet.objects.create(
            titre="Projet Prototype Agri",
            domaine=self.domaine_agri,
            cohorte=self.cohorte_2025,
            porteur=self.porteur_homme,
            statut=Projet.Statut.PROTOTYPE,
        )

    def test_indicateurs_endpoint_contract_and_values(self):
        """GET /api/analytics/indicateurs/ renvoie le contrat exact avec calculs justes."""
        url = reverse("analytics-indicateurs")
        res = self.client.get(url)
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.json()

        # Clés requises
        required_keys = {
            "idees_recues",
            "dossiers_deposes",
            "projets_selectionnes",
            "projets_incubes",
            "prototypes",
            "taux_selection",
        }
        self.assertEqual(set(data.keys()), required_keys)

        # Vérification des valeurs :
        # Total idées = 5
        # Déposés (statuts != idee) = 4
        # Sélectionnés (selectionne + incube + prototype) = 3
        # Incubés (incube + prototype) = 2
        # Prototypes = 1
        # Taux de sélection = 3 / 4 * 100 = 75.0 %
        self.assertEqual(data["idees_recues"], 5)
        self.assertEqual(data["dossiers_deposes"], 4)
        self.assertEqual(data["projets_selectionnes"], 3)
        self.assertEqual(data["projets_incubes"], 2)
        self.assertEqual(data["prototypes"], 1)
        self.assertEqual(data["taux_selection"], 75.0)

    def test_par_domaine_endpoint(self):
        """GET /api/analytics/par-domaine/ regroupe les projets par domaine."""
        url = reverse("analytics-par-domaine")
        res = self.client.get(url)
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.json()

        self.assertIsInstance(data, list)
        self.assertGreater(len(data), 0)
        # Format d'élément
        for item in data:
            self.assertIn("domaine", item)
            self.assertIn("total", item)

        # TIC a 3 projets, Agri en a 2
        tic_item = next((item for item in data if "TIC" in item["domaine"]), None)
        agri_item = next((item for item in data if "Agriculture" in item["domaine"]), None)
        self.assertIsNotNone(tic_item)
        self.assertIsNotNone(agri_item)
        self.assertEqual(tic_item["total"], 3)
        self.assertEqual(agri_item["total"], 2)

    def test_par_cohorte_endpoint(self):
        """GET /api/analytics/par-cohorte/ compare les cohortes selon le contrat."""
        url = reverse("analytics-par-cohorte")
        res = self.client.get(url)
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.json()

        self.assertIsInstance(data, list)
        for item in data:
            self.assertIn("cohorte", item)
            self.assertIn("total", item)
            self.assertIn("selectionnes", item)
            self.assertIn("prototypes", item)

        c2024 = next((item for item in data if item["cohorte"] == "Cohorte 2024"), None)
        c2025 = next((item for item in data if item["cohorte"] == "Cohorte 2025"), None)
        self.assertIsNotNone(c2024)
        self.assertIsNotNone(c2025)
        self.assertEqual(c2024["total"], 2)
        self.assertEqual(c2025["total"], 3)
        self.assertEqual(c2025["selectionnes"], 3)
        self.assertEqual(c2025["prototypes"], 1)

    def test_par_genre_endpoint(self):
        """GET /api/analytics/par-genre/ renvoie la répartition Femmes, Hommes, Autre."""
        url = reverse("analytics-par-genre")
        res = self.client.get(url)
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.json()

        self.assertIsInstance(data, list)
        genre_map = {item["genre"]: item["total"] for item in data}
        self.assertIn("Femmes", genre_map)
        self.assertIn("Hommes", genre_map)
        self.assertEqual(genre_map["Femmes"], 2)  # 2 projets portés par femme
        self.assertEqual(genre_map["Hommes"], 2)  # 2 projets portés par homme
        self.assertEqual(genre_map["Autre / non précisé"], 1)

    def test_evolution_endpoint(self):
        """GET /api/analytics/evolution/ renvoie l'évolution temporelle au format YYYY-MM."""
        url = reverse("analytics-evolution")
        res = self.client.get(url)
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.json()

        self.assertIsInstance(data, list)
        current_month = date.today().strftime("%Y-%m")
        # Les 5 projets créés aujourd'hui doivent se trouver dans le mois courant
        month_item = next((item for item in data if item["mois"] == current_month), None)
        self.assertIsNotNone(month_item)
        self.assertEqual(month_item["total"], 5)

    def test_filters_combination(self):
        """Les filtres ?cohorte= &domaine= &statut= &debut= &fin= se combinent fidèlement."""
        url = reverse("analytics-indicateurs")

        # Filtre sur cohorte 2024
        res_cohorte = self.client.get(url, {"cohorte": "2024"})
        self.assertEqual(res_cohorte.status_code, status.HTTP_200_OK)
        self.assertEqual(res_cohorte.json()["idees_recues"], 2)

        # Filtre sur domaine TIC
        res_domaine = self.client.get(url, {"domaine": "tic-numerique"})
        self.assertEqual(res_domaine.status_code, status.HTTP_200_OK)
        self.assertEqual(res_domaine.json()["idees_recues"], 3)

        # Filtre combiné : Cohorte 2024 + TIC
        res_combine = self.client.get(url, {"cohorte": "2024", "domaine": "tic-numerique"})
        self.assertEqual(res_combine.status_code, status.HTTP_200_OK)
        self.assertEqual(res_combine.json()["idees_recues"], 1)

        # Filtre statut
        res_statut = self.client.get(url, {"statut": "prototype"})
        self.assertEqual(res_statut.status_code, status.HTTP_200_OK)
        self.assertEqual(res_statut.json()["idees_recues"], 1)
        self.assertEqual(res_statut.json()["prototypes"], 1)

    def test_empty_dataset_edge_case(self):
        """Quand aucun projet ne correspond aux filtres, pas d'erreur 500 ni division par zéro."""
        url = reverse("analytics-indicateurs")
        res = self.client.get(url, {"cohorte": "9999"})  # cohorte inexistante
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.json()
        self.assertEqual(data["idees_recues"], 0)
        self.assertEqual(data["dossiers_deposes"], 0)
        self.assertEqual(data["taux_selection"], 0.0)


# ==============================================================================
# TESTS T3 — Export Excel et PDF du Tableau de Bord Décisionnel
# ==============================================================================

@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class AnalyticsExportTests(TestCase):
    """Vérification des exports XLSX et PDF avec respect de la charte CIAT et des filtres."""

    def setUp(self):
        self.client = APIClient()

        self.coordo = User.objects.create_user(
            email="coordo.export@ciat.sn",
            password="pwd",
            role=Role.COORDONNATEUR,
        )
        self.admin = User.objects.create_user(
            email="admin.export@ciat.sn",
            password="pwd",
            role=Role.ADMIN,
        )
        self.comite = User.objects.create_user(
            email="comite.export@ciat.sn",
            password="pwd",
            role=Role.COMITE,
        )
        self.porteur = User.objects.create_user(
            email="porteur.export@ciat.sn",
            password="pwd",
            role=Role.PORTEUR,
        )

        self.domaine_tic = Domaine.objects.create(nom="TIC & Numérique", slug="tic-numerique")
        self.domaine_agri = Domaine.objects.create(nom="AgriTech", slug="agritech")
        self.cohorte_2025 = Cohorte.objects.create(nom="Cohorte 2025", annee=2025)

        # Création de quelques projets avec caractères accentués
        self.p1 = Projet.objects.create(
            titre="Système d'irrigation intelligente à énergie solaire",
            resume="Projet innovant pour les agriculteurs sénégalais.",
            domaine=self.domaine_agri,
            cohorte=self.cohorte_2025,
            porteur=self.porteur,
            statut=Projet.Statut.INCUBE,
        )
        self.p2 = Projet.objects.create(
            titre="Plateforme médicale de télémédecine rurale",
            resume="Consultation à distance.",
            domaine=self.domaine_tic,
            cohorte=self.cohorte_2025,
            porteur=self.porteur,
            statut=Projet.Statut.PROTOTYPE,
        )

    def test_export_xlsx_valide(self):
        """GET /api/analytics/export/?format=xlsx génère un fichier Excel valide avec la charte CIAT."""
        self.client.force_authenticate(user=self.coordo)
        res = self.client.get(reverse("analytics-export"), {"format": "xlsx"})

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(
            res["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment; filename=", res["Content-Disposition"])
        self.assertTrue(res["Content-Disposition"].endswith('.xlsx"'))

        # Lecture du fichier via openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        self.assertIn("Analyses CIAT", wb.sheetnames)
        ws = wb["Analyses CIAT"]

        # Vérification de l'en-tête institutionnel et des accents
        self.assertIn("CENTRE D'INNOVATION", str(ws["A1"].value))
        self.assertIn("Tableau de Bord", str(ws["A2"].value))

        # Vérification qu'il y a des lignes de données
        textes_cellules = [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value]
        self.assertTrue(any("INDICATEURS CLÉS" in t for t in textes_cellules))
        self.assertTrue(any("AgriTech" in t for t in textes_cellules))

    def test_export_pdf_valide(self):
        """GET /api/analytics/export/?format=pdf génère un fichier PDF valide avec pagination et charte."""
        self.client.force_authenticate(user=self.admin)
        res = self.client.get(reverse("analytics-export"), {"format": "pdf"})

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res["Content-Type"], "application/pdf")
        self.assertIn("attachment; filename=", res["Content-Disposition"])
        self.assertTrue(res["Content-Disposition"].endswith('.pdf"'))

        # Header magic bytes PDF
        self.assertTrue(res.content.startswith(b"%PDF"))
        self.assertGreater(len(res.content), 1000)

    def test_export_avec_filtres_actifs(self):
        """Les filtres (domaine, cohorte, statut) s'appliquent correctement aux exports."""
        self.client.force_authenticate(user=self.coordo)
        res = self.client.get(
            reverse("analytics-export"),
            {"format": "xlsx", "domaine": "agritech", "statut": "incube"},
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)

        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb["Analyses CIAT"]
        # Vérifie que la ligne de métadonnées mentionne le filtre
        self.assertIn("agritech", str(ws["A3"].value))

    def test_export_format_invalide_ou_manquant_renvoie_400(self):
        """Un format inconnu ou manquant renvoie 400 Bad Request."""
        self.client.force_authenticate(user=self.coordo)

        # Format manquant
        res_manquant = self.client.get(reverse("analytics-export"))
        self.assertEqual(res_manquant.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("Format d'export invalide", res_manquant.json()["detail"])

        # Format invalide
        res_invalide = self.client.get(reverse("analytics-export"), {"format": "csv"})
        self.assertEqual(res_invalide.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("xlsx", res_invalide.json()["detail"])

    def test_export_refuse_au_porteur(self):
        """Un porteur de projet n'a pas accès à l'export des analyses (403)."""
        self.client.force_authenticate(user=self.porteur)
        res = self.client.get(reverse("analytics-export"), {"format": "xlsx"})
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

    def test_export_autorise_au_comite(self):
        """Un membre du comité de sélection a accès à l'export (200)."""
        self.client.force_authenticate(user=self.comite)
        res = self.client.get(reverse("analytics-export"), {"format": "xlsx"})
        self.assertEqual(res.status_code, status.HTTP_200_OK)

    def test_export_refuse_aux_anonymes(self):
        """Un visiteur non connecté est rejeté (401)."""
        res = self.client.get(reverse("analytics-export"), {"format": "xlsx"})
        self.assertEqual(res.status_code, status.HTTP_401_UNAUTHORIZED)


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class AnalyticsAIQuestionTests(TestCase):
    """Tests complets pour l'interrogation en langage naturel (T2)."""

    def setUp(self):
        self.client = APIClient()
        self.url = reverse("analytics-question")

        # Utilisateurs
        self.admin = User.objects.create_user(
            email="admin.ai@ciat.sn", password="pwd", role=Role.ADMIN
        )
        self.coordo = User.objects.create_user(
            email="coordo.ai@ciat.sn", password="pwd", role=Role.COORDONNATEUR
        )
        self.comite = User.objects.create_user(
            email="comite.ai@ciat.sn", password="pwd", role=Role.COMITE
        )
        self.porteur = User.objects.create_user(
            email="porteur.ai@ciat.sn", password="pwd", role=Role.PORTEUR
        )
        self.visiteur = User.objects.create_user(
            email="visiteur.ai@ciat.sn", password="pwd", role=Role.VISITEUR
        )

        # Domaines et cohortes
        self.domaine_tic = Domaine.objects.create(nom="TIC & Numérique", slug="tic")
        self.domaine_agri = Domaine.objects.create(nom="Agroécologie", slug="agri")
        self.cohorte_2024 = Cohorte.objects.create(nom="Cohorte 2024", annee=2024)
        self.cohorte_2025 = Cohorte.objects.create(nom="Cohorte 2025", annee=2025)

        # Porteurs avec genre
        self.porteur_f = User.objects.create_user(
            email="femme.ai@ciat.sn", password="pwd", role=Role.PORTEUR, genre="F"
        )
        self.porteur_h = User.objects.create_user(
            email="homme.ai@ciat.sn", password="pwd", role=Role.PORTEUR, genre="H"
        )

        # Projets
        Projet.objects.create(
            titre="Solution IoT Smart Agriculture",
            resume="Projet IoT",
            domaine=self.domaine_agri,
            cohorte=self.cohorte_2025,
            porteur=self.porteur_f,
            statut=Projet.Statut.PROTOTYPE,
        )
        Projet.objects.create(
            titre="Plateforme E-santé",
            resume="Santé digitale",
            domaine=self.domaine_tic,
            cohorte=self.cohorte_2025,
            porteur=self.porteur_h,
            statut=Projet.Statut.INCUBE,
        )
        Projet.objects.create(
            titre="Ancien projet 2024",
            resume="Archive 2024",
            domaine=self.domaine_tic,
            cohorte=self.cohorte_2024,
            porteur=self.porteur_f,
            statut=Projet.Statut.DEPOSE,
        )

    def test_anonymous_and_unauthorized_roles_rejected(self):
        """Les requêtes anonymes ou de porteurs/visiteurs sont bloquées (401/403)."""
        # Anonyme
        res_anon = self.client.post(self.url, {"question": "Combien de projets ?"})
        self.assertIn(
            res_anon.status_code,
            [status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN],
        )

        # Porteur
        self.client.force_authenticate(user=self.porteur)
        res_porteur = self.client.post(self.url, {"question": "Combien de projets ?"})
        self.assertEqual(res_porteur.status_code, status.HTTP_403_FORBIDDEN)

        # Visiteur
        self.client.force_authenticate(user=self.visiteur)
        res_visiteur = self.client.post(self.url, {"question": "Combien de projets ?"})
        self.assertEqual(res_visiteur.status_code, status.HTTP_403_FORBIDDEN)

    def test_empty_or_missing_question_returns_400(self):
        """Une question absente ou vide renvoie HTTP 400 Bad Request."""
        self.client.force_authenticate(user=self.coordo)

        res_empty = self.client.post(self.url, {})
        self.assertEqual(res_empty.status_code, status.HTTP_400_BAD_REQUEST)

        res_whitespace = self.client.post(self.url, {"question": "    "})
        self.assertEqual(res_whitespace.status_code, status.HTTP_400_BAD_REQUEST)

    @patch("apps.analytics.services.ai.chat")
    def test_prototype_in_cohorte_2025_question_success(self, mock_chat):
        """Vérifie : 'Combien de projets en prototype dans la cohorte 2025 ?' donne le bon chiffre."""
        import json
        self.client.force_authenticate(user=self.coordo)

        # Mocking : 1er appel = routeur JSON, 2e appel = synthèse
        router_response = json.dumps({
            "endpoint": "indicateurs",
            "filtres": {"cohorte": "2025", "statut": "prototype"},
            "hors_sujet": False,
        })
        synthesis_response = "Il y a exactement 1 projet au stade de prototype dans la cohorte 2025."
        mock_chat.side_effect = [router_response, synthesis_response]

        res = self.client.post(
            self.url,
            {"question": "Combien de projets en prototype dans la cohorte 2025 ?"},
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.json()

        self.assertIn("reponse", data)
        self.assertIn("donnees", data)
        self.assertIn("filtres_utilises", data)

        self.assertEqual(data["reponse"], synthesis_response)
        self.assertEqual(data["filtres_utilises"]["cohorte"], "2025")
        self.assertEqual(data["filtres_utilises"]["statut"], "prototype")

        # Recoupement des données ORM exactes calculées par Django
        self.assertEqual(data["donnees"]["prototypes"], 1)
        self.assertEqual(data["donnees"]["idees_recues"], 1)

    @patch("apps.analytics.services.ai.chat")
    def test_par_domaine_question(self, mock_chat):
        """Vérifie le routage vers l'endpoint 'par-domaine' et les données agrégées."""
        import json
        self.client.force_authenticate(user=self.admin)

        router_response = json.dumps({
            "endpoint": "par-domaine",
            "filtres": {},
            "hors_sujet": False,
        })
        synthesis_response = "Les projets se répartissent entre TIC (2 projets) et Agroécologie (1 projet)."
        mock_chat.side_effect = [router_response, synthesis_response]

        res = self.client.post(
            self.url,
            {"question": "Quelle est la répartition des projets par domaine ?"},
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.json()
        self.assertEqual(data["reponse"], synthesis_response)
        self.assertIsInstance(data["donnees"], list)
        domaines = {d["domaine"]: d["total"] for d in data["donnees"]}
        self.assertEqual(domaines["TIC & Numérique"], 2)
        self.assertEqual(domaines["Agroécologie"], 1)

    @patch("apps.analytics.services.ai.chat")
    def test_par_genre_question(self, mock_chat):
        """Vérifie le routage vers l'endpoint 'par-genre'."""
        import json
        self.client.force_authenticate(user=self.comite)

        router_response = json.dumps({
            "endpoint": "par-genre",
            "filtres": {},
            "hors_sujet": False,
        })
        synthesis_response = "2 projets sont portés par des femmes et 1 par un homme."
        mock_chat.side_effect = [router_response, synthesis_response]

        res = self.client.post(
            self.url,
            {"question": "Quelle est la part des femmes parmi les porteurs de projet ?"},
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.json()
        genres = {g["genre"]: g["total"] for g in data["donnees"]}
        self.assertEqual(genres["Femmes"], 2)
        self.assertEqual(genres["Hommes"], 1)

    @patch("apps.analytics.services.ai.chat")
    def test_off_topic_question_is_refused(self, mock_chat):
        """Une question hors sujet est explicitement refusée selon les consignes."""
        import json
        self.client.force_authenticate(user=self.coordo)

        router_response = json.dumps({
            "endpoint": "hors_sujet",
            "filtres": {},
            "hors_sujet": True,
        })
        mock_chat.return_value = router_response

        res = self.client.post(
            self.url,
            {"question": "Quel temps fait-il aujourd'hui à Diamniadio ?"},
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.json()
        self.assertEqual(
            data["reponse"],
            "Je ne peux pas répondre à cette question avec les données disponibles.",
        )
        self.assertIsNone(data["donnees"])
        self.assertEqual(data["filtres_utilises"], {})

    @patch("apps.analytics.services.ai.chat")
    def test_security_injection_endpoint_refused(self, mock_chat):
        """Toute tentative d'injection d'endpoint non autorisé est rejetée."""
        import json
        self.client.force_authenticate(user=self.coordo)

        # Tentative d'injection d'un endpoint arbitraire
        router_response = json.dumps({
            "endpoint": "DROP_TABLE_PROJETS",
            "filtres": {"cohorte": "2025"},
            "hors_sujet": False,
        })
        mock_chat.return_value = router_response

        res = self.client.post(
            self.url,
            {"question": "Supprime les projets"},
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.json()
        self.assertEqual(
            data["reponse"],
            "Je ne peux pas répondre à cette question avec les données disponibles.",
        )
        self.assertIsNone(data["donnees"])

    @patch("apps.analytics.services.ai.chat")
    def test_ai_not_configured_returns_503(self, mock_chat):
        """Quand aucune clé d'API n'est configurée, l'API renvoie 503 Service Unavailable."""
        from apps.core.ai import AINotConfigured
        self.client.force_authenticate(user=self.coordo)

        mock_chat.side_effect = AINotConfigured("Aucune clé d'API IA")

        res = self.client.post(
            self.url,
            {"question": "Combien de projets ?"},
        )
        self.assertEqual(res.status_code, status.HTTP_503_SERVICE_UNAVAILABLE)
        self.assertIn("Le service d'analyse IA n'est pas configuré", res.json().get("detail", ""))

    @patch("apps.analytics.services.ai.chat")
    def test_ai_error_returns_503(self, mock_chat):
        """Quand les fournisseurs IA échouent (quota, timeout), l'API renvoie 503."""
        from apps.core.ai import AIError
        self.client.force_authenticate(user=self.coordo)

        mock_chat.side_effect = AIError("Quota atteint")

        res = self.client.post(
            self.url,
            {"question": "Combien de projets ?"},
        )
        self.assertEqual(res.status_code, status.HTTP_503_SERVICE_UNAVAILABLE)
        self.assertIn("Le service d'analyse IA est momentanément indisponible", res.json().get("detail", ""))


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class AnalyticsPartenariatsAndFinancementsModelsTests(TestCase):
    """Validation des règles métier strictes des modèles Partenariat et Financement (T3)."""

    def setUp(self):
        self.domaine = Domaine.objects.create(nom="TIC", slug="tic")
        self.cohorte = Cohorte.objects.create(nom="Cohorte 2025", annee=2025)
        self.porteur = User.objects.create_user(
            email="porteur.t3@ciat.sn", password="pwd", role=Role.PORTEUR
        )

        self.projet_idee = Projet.objects.create(
            titre="Idée sans partenariat",
            resume="Test idée",
            domaine=self.domaine,
            cohorte=self.cohorte,
            porteur=self.porteur,
            statut=Projet.Statut.IDEE,
        )

        self.projet_depose = Projet.objects.create(
            titre="Dépôt sans financement",
            resume="Test dépose",
            domaine=self.domaine,
            cohorte=self.cohorte,
            porteur=self.porteur,
            statut=Projet.Statut.DEPOSE,
        )

        self.projet_incube = Projet.objects.create(
            titre="Projet Incubé Éligible",
            resume="Incubation",
            domaine=self.domaine,
            cohorte=self.cohorte,
            porteur=self.porteur,
            statut=Projet.Statut.INCUBE,
        )

    def test_partenariat_creation_allowed_for_incubated_project(self):
        """Un projet sélectionné ou incubé peut recevoir un partenariat."""
        part = Partenariat.objects.create(
            projet=self.projet_incube,
            partenaire="Sonatel",
            type=TypePartenariat.TECHNIQUE,
            date_debut=date.today(),
            statut=StatutPartenariat.ACTIF,
        )
        self.assertIsNotNone(part.id)
        self.assertEqual(part.partenaire, "Sonatel")

    def test_partenariat_forbidden_for_idea_or_depose_project(self):
        """Une ValidationError est levée si on tente d'associer un partenariat à une simple idée."""
        with self.assertRaises(ValidationError):
            Partenariat.objects.create(
                projet=self.projet_idee,
                partenaire="DER/FJ",
                type=TypePartenariat.FINANCIER,
                date_debut=date.today(),
            )

        with self.assertRaises(ValidationError):
            Partenariat.objects.create(
                projet=self.projet_depose,
                partenaire="DER/FJ",
                type=TypePartenariat.FINANCIER,
                date_debut=date.today(),
            )

    def test_financement_creation_allowed_for_incubated_project(self):
        """Un projet incubé peut recevoir un financement."""
        fin = Financement.objects.create(
            projet=self.projet_incube,
            source="DER/FJ",
            montant=Decimal("5000000"),
            date=date.today(),
            statut=StatutFinancement.DECAISSE,
        )
        self.assertIsNotNone(fin.id)
        self.assertEqual(fin.montant, Decimal("5000000"))

    def test_financement_forbidden_for_idea_or_depose_project(self):
        """Une ValidationError est levée si on tente d'associer un financement à une idée ou dossier déposé."""
        with self.assertRaises(ValidationError):
            Financement.objects.create(
                projet=self.projet_idee,
                source="Fonds Amorçage",
                montant=Decimal("3000000"),
                date=date.today(),
            )

        with self.assertRaises(ValidationError):
            Financement.objects.create(
                projet=self.projet_depose,
                source="Fonds Amorçage",
                montant=Decimal("3000000"),
                date=date.today(),
            )

    def test_financement_negative_amount_raises_validation_error(self):
        """Un montant négatif ou nul est strictement interdit."""
        with self.assertRaises(ValidationError):
            Financement.objects.create(
                projet=self.projet_incube,
                source="DER/FJ",
                montant=Decimal("-500000"),
                date=date.today(),
            )


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class AnalyticsPartenariatsAndFinancementsEndpointsTests(TestCase):
    """Vérification des endpoints DRF pour les partenariats et financements (T3)."""

    def setUp(self):
        self.client = APIClient()

        # Utilisateurs
        self.admin = User.objects.create_user(
            email="admin.pf@ciat.sn", password="pwd", role=Role.ADMIN
        )
        self.coordo = User.objects.create_user(
            email="coordo.pf@ciat.sn", password="pwd", role=Role.COORDONNATEUR
        )
        self.comite = User.objects.create_user(
            email="comite.pf@ciat.sn", password="pwd", role=Role.COMITE
        )
        self.porteur = User.objects.create_user(
            email="porteur.pf@ciat.sn", password="pwd", role=Role.PORTEUR
        )
        self.visiteur = User.objects.create_user(
            email="visiteur.pf@ciat.sn", password="pwd", role=Role.VISITEUR
        )

        # Domaines et cohortes
        self.domaine_tic = Domaine.objects.create(nom="TIC Numérique", slug="tic")
        self.domaine_agri = Domaine.objects.create(nom="Agroécologie", slug="agri")
        self.cohorte_2024 = Cohorte.objects.create(nom="Cohorte 2024", annee=2024)
        self.cohorte_2025 = Cohorte.objects.create(nom="Cohorte 2025", annee=2025)

        # Projets éligibles
        self.projet_1 = Projet.objects.create(
            titre="Smart Agri IoT",
            resume="IoT Agri",
            domaine=self.domaine_agri,
            cohorte=self.cohorte_2025,
            porteur=self.porteur,
            statut=Projet.Statut.PROTOTYPE,
        )
        self.projet_2 = Projet.objects.create(
            titre="E-Santé Dakar",
            resume="Santé digitale",
            domaine=self.domaine_tic,
            cohorte=self.cohorte_2024,
            porteur=self.porteur,
            statut=Projet.Statut.INCUBE,
        )

        # Partenariats
        self.part_1 = Partenariat.objects.create(
            projet=self.projet_1,
            partenaire="DER/FJ",
            type=TypePartenariat.FINANCIER,
            date_debut=date.today(),
            statut=StatutPartenariat.ACTIF,
        )
        self.part_2 = Partenariat.objects.create(
            projet=self.projet_2,
            partenaire="Orange Fab",
            type=TypePartenariat.COMMERCIAL,
            date_debut=date.today() - timedelta(days=60),
            statut=StatutPartenariat.SIGNE,
        )

        # Financements
        self.fin_1 = Financement.objects.create(
            projet=self.projet_1,
            source="DER/FJ",
            montant=Decimal("10000000"),
            date=date.today(),
            statut=StatutFinancement.DECAISSE,
        )
        self.fin_2 = Financement.objects.create(
            projet=self.projet_2,
            source="CIAT Amorçage",
            montant=Decimal("5000000"),
            date=date.today() - timedelta(days=30),
            statut=StatutFinancement.ACCORDE,
        )

        self.url_part = reverse("analytics-partenariats")
        self.url_fin = reverse("analytics-financements")
        self.url_stats = reverse("analytics-partenariats-financements-stats")

    def test_rbac_permissions_on_endpoints(self):
        """Seuls l'équipe CIAT et le comité peuvent consulter les partenariats et financements."""
        # Anonyme
        for url in [self.url_part, self.url_fin, self.url_stats]:
            res = self.client.get(url)
            self.assertIn(
                res.status_code,
                [status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN],
            )

        # Porteur
        self.client.force_authenticate(user=self.porteur)
        for url in [self.url_part, self.url_fin, self.url_stats]:
            res = self.client.get(url)
            self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

        # Visiteur
        self.client.force_authenticate(user=self.visiteur)
        for url in [self.url_part, self.url_fin, self.url_stats]:
            res = self.client.get(url)
            self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

        # Rôles autorisés
        for u in [self.coordo, self.admin, self.comite]:
            self.client.force_authenticate(user=u)
            for url in [self.url_part, self.url_fin, self.url_stats]:
                res = self.client.get(url)
                self.assertEqual(res.status_code, status.HTTP_200_OK)

    def test_partenariats_list_and_filters(self):
        """GET /api/analytics/partenariats/ filtre fidèlement par cohorte, domaine et statut."""
        self.client.force_authenticate(user=self.coordo)

        res = self.client.get(self.url_part)
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.json()
        self.assertEqual(len(data), 2)

        # Filtre sur cohorte 2025
        res_c25 = self.client.get(self.url_part, {"cohorte": "2025"})
        self.assertEqual(res_c25.status_code, status.HTTP_200_OK)
        data_c25 = res_c25.json()
        self.assertEqual(len(data_c25), 1)
        self.assertEqual(data_c25[0]["partenaire"], "DER/FJ")

        # Filtre sur type
        res_type = self.client.get(self.url_part, {"type": "commercial"})
        self.assertEqual(res_type.status_code, status.HTTP_200_OK)
        data_type = res_type.json()
        self.assertEqual(len(data_type), 1)
        self.assertEqual(data_type[0]["partenaire"], "Orange Fab")

    def test_financements_list_and_filters(self):
        """GET /api/analytics/financements/ filtre fidèlement par source et domaine."""
        self.client.force_authenticate(user=self.coordo)

        res = self.client.get(self.url_fin)
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.json()
        self.assertEqual(len(data), 2)

        # Filtre sur source DER/FJ
        res_src = self.client.get(self.url_fin, {"source": "DER/FJ"})
        self.assertEqual(res_src.status_code, status.HTTP_200_OK)
        self.assertEqual(len(res_src.json()), 1)

        # Filtre sur domaine agri
        res_dom = self.client.get(self.url_fin, {"domaine": "agri"})
        self.assertEqual(res_dom.status_code, status.HTTP_200_OK)
        self.assertEqual(len(res_dom.json()), 1)

    def test_partenariats_financements_stats(self):
        """GET /api/analytics/partenariats-financements/stats/ calcule exactement les indicateurs."""
        self.client.force_authenticate(user=self.coordo)

        res = self.client.get(self.url_stats)
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        stats_data = res.json()

        # Total financé = 10 000 000 + 5 000 000 = 15 000 000
        # Décaissé = 10 000 000
        # Partenariats actifs = 1 (part_1)
        # Total partenariats = 2
        self.assertEqual(stats_data["total_finance"], 15000000)
        self.assertEqual(stats_data["financements_decaisses"], 10000000)
        self.assertEqual(stats_data["nombre_financements"], 2)
        self.assertEqual(stats_data["nombre_partenariats_actifs"], 1)
        self.assertEqual(stats_data["total_partenariats"], 2)
        self.assertEqual(len(stats_data["par_source"]), 2)


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class AnalyticsSeedAnalyticsCommandTests(TestCase):
    """Vérification de la commande de remplissage seed_analytics (idempotence et cohérence)."""

    def setUp(self):
        self.domaine = Domaine.objects.create(nom="TIC Test", slug="tic-test")
        self.cohorte = Cohorte.objects.create(nom="Cohorte Test 2025", annee=2025)
        self.porteur = User.objects.create_user(
            email="porteur.seed@ciat.sn", password="pwd", role=Role.PORTEUR
        )

        # 1 Idée (ne doit JAMAIS recevoir de financement)
        Projet.objects.create(
            titre="Idée pure",
            resume="Idée",
            domaine=self.domaine,
            cohorte=self.cohorte,
            porteur=self.porteur,
            statut=Projet.Statut.IDEE,
        )

        # 2 Projets éligibles (Incubé & Prototype)
        Projet.objects.create(
            titre="Projet Éligible 1",
            resume="Incubé",
            domaine=self.domaine,
            cohorte=self.cohorte,
            porteur=self.porteur,
            statut=Projet.Statut.INCUBE,
        )
        Projet.objects.create(
            titre="Projet Éligible 2",
            resume="Prototype",
            domaine=self.domaine,
            cohorte=self.cohorte,
            porteur=self.porteur,
            statut=Projet.Statut.PROTOTYPE,
        )

    def test_seed_analytics_is_idempotent_and_respects_rules(self):
        """seed_analytics génère des données valides et ne crée aucun doublon au second passage."""
        # Premier passage
        call_command("seed_analytics")
        nb_part_1 = Partenariat.objects.count()
        nb_fin_1 = Financement.objects.count()

        self.assertGreater(nb_part_1, 0)
        self.assertGreater(nb_fin_1, 0)

        # Règle d'or : aucune idée ni dossier déposé n'a de financement
        self.assertEqual(
            Financement.objects.filter(projet__statut=Projet.Statut.IDEE).count(), 0
        )
        self.assertEqual(
            Partenariat.objects.filter(projet__statut=Projet.Statut.IDEE).count(), 0
        )

        # Deuxième passage (idempotence)
        call_command("seed_analytics")
        nb_part_2 = Partenariat.objects.count()
        nb_fin_2 = Financement.objects.count()

        self.assertEqual(nb_part_1, nb_part_2)
        self.assertEqual(nb_fin_1, nb_fin_2)


